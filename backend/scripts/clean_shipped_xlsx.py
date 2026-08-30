"""
Strip Beaufort/Batch-92 real data from the shipped Feed Program xlsx files
so new users see a clean template instead of Jason's actual farm data.

Rules:
  - KEEP: all formulas (they're template calculations)
  - KEEP: all strings (headers/labels)
  - KEEP: "Consumption Guide" sheet entirely (static Ross 308 breed reference)
  - KEEP: all styles/formatting/column widths/merged cells
  - CLEAR: numbers and dates in user-data sheets (SHED *, Beaufort *, Pickup *, end of batch)

Also renames the "Beaufort 86" tab to "Current Batch" so no one sees the grower's name.

Run once from /app/backend to regenerate the clean templates.
"""
import shutil
from pathlib import Path

import openpyxl

FEED_PROGRAM_XLSX = Path("/app/silo/artifacts/feed-program/public/feed-program.xlsx")
BATCH_RESULTS_XLSX = Path("/app/silo/artifacts/feed-program/public/batch-results.xlsx")

# Sheets that should have their user-data cells wiped (numbers + dates only,
# formulas + strings + styles stay). Match by name-startswith so "Beaufort 86",
# "Beaufort 87", "SHED 1 & 2", etc. all get caught.
USER_DATA_PREFIXES = ("SHED", "Beaufort", "sheds CFCR", "Pickup sheet",
                      "Pickup Sheet", "end of batch")

# Sheets to fully preserve (no clearing at all)
PRESERVE_ENTIRELY = ("Consumption Guide", "WEEKLY STOCK TAKE")

# "Beaufort 86" is a batch tab that displays as a grower-facing name. Rename
# to something generic so no new user sees "Beaufort" as the default.
RENAME_MAP = {
    "Beaufort 86": "Current Batch",
    "Beaufort 87": "Current Batch",
    "Beaufort 88": "Current Batch",
}


def is_user_data_sheet(name: str) -> bool:
    if name.strip() in PRESERVE_ENTIRELY:
        return False
    return any(name.strip().startswith(p) for p in USER_DATA_PREFIXES)


def clean_workbook(path: Path) -> tuple[int, int]:
    """Return (cells_cleared, sheets_renamed)."""
    wb = openpyxl.load_workbook(path, data_only=False)
    cleared = 0
    renamed = 0

    for name in list(wb.sheetnames):
        ws = wb[name]
        stripped = name.strip()

        # Rename Beaufort tabs to generic name
        if stripped in RENAME_MAP:
            new_name = RENAME_MAP[stripped]
            # Avoid dupes across the workbook
            if new_name not in wb.sheetnames:
                ws.title = new_name
                renamed += 1

        if not is_user_data_sheet(name):
            continue

        # Iterate every cell — clear number/date values, preserve formulas + strings + styles
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                # Formulas start with "=" — always keep (template logic)
                if isinstance(v, str) and v.startswith("="):
                    continue
                # Strings are headers/labels — keep
                if isinstance(v, str):
                    continue
                # Numbers + dates + booleans → clear (this is the leaked data)
                cell.value = None
                cleared += 1

    wb.save(path)
    return cleared, renamed


def main() -> None:
    for src in (FEED_PROGRAM_XLSX, BATCH_RESULTS_XLSX):
        if not src.exists():
            print(f"  ✗ missing: {src}")
            continue
        backup = src.with_suffix(f".pre-clean-backup.xlsx")
        if not backup.exists():
            shutil.copy2(src, backup)
            print(f"  ↳ backup saved: {backup.name}")
        cleared, renamed = clean_workbook(src)
        print(f"  ✓ {src.name}: cleared {cleared} data cells, renamed {renamed} tabs")


if __name__ == "__main__":
    main()
